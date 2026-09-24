import io
import os
from uuid import UUID

import pytest
from docx import Document
from openpyxl import Workbook
from test_platform_postgres import platform as _platform_fixture

from agent_platform.modules.knowledge.chunking import ChunkingPolicy, chunk_pages
from agent_platform.modules.knowledge.parsers import parse_file
from agent_platform.modules.knowledge.retrieval import fuse

platform = _platform_fixture

ROOT = "/api/v1/knowledge-workspace"


@pytest.mark.asyncio
@pytest.mark.postgres
async def test_semantic_preview_cache_reused_for_pages_and_save(platform):
    import asyncio
    from unittest.mock import AsyncMock

    from test_model_gateway import setup

    client, services, settings = platform
    _, _, profile = await setup(client, "embed")
    w = services.platform.knowledge_workspace
    upload = await client.post(
        ROOT + "/uploads?filename=semantic.md", content=("知识段落。" * 500).encode()
    )
    did = upload.json()["id"]
    await w.run_next(settings.tenant_id)
    policy = {
        "strategy": "semantic",
        "semantic_profile_id": profile["id"],
        "semantic_profile_version": 1,
    }
    original = w.gateway.invoke
    w.gateway.invoke = AsyncMock(wraps=original)
    try:
        path = ROOT + f"/documents/{did}"
        first, second = await asyncio.gather(
            client.post(path + "/preview?limit=1", json=policy),
            client.post(path + "/preview?limit=1&offset=1", json=policy),
        )
        assert first.status_code == second.status_code == 200
        calls = w.gateway.invoke.call_count
        assert calls > 0
        await client.post(path + "/preview?limit=1", json=policy)
        assert w.gateway.invoke.call_count == calls
        response = await client.put(
            path + "/draft",
            json={
                "revision": 1,
                "policy": policy,
                "offset": 0,
                "replace_count": 1,
                "regenerate": True,
                "chunks": [{"content": "编辑后的语义分片"}],
            },
        )
        assert response.status_code == 200, response.text
        assert w.gateway.invoke.call_count == calls
        changed = await client.post(path + "/preview", json={**policy, "semantic_threshold": 0.9})
        assert changed.status_code == 200
        assert w.gateway.invoke.call_count > calls
    finally:
        w.gateway.invoke = original


@pytest.mark.asyncio
@pytest.mark.postgres
async def test_paginated_chunks_preview_and_partial_save(platform):
    client, services, settings = platform
    content = "\n\n".join(f"Section {i}: " + "content " * 100 for i in range(65))
    upload = await client.post(ROOT + "/uploads?filename=pagination.md", content=content.encode())
    did = upload.json()["id"]
    await services.platform.knowledge_workspace.run_next(settings.tenant_id)
    path = ROOT + f"/documents/{did}"
    first = (await client.get(path)).json()
    assert len(first["chunks"]) == 20 and first["total"] > 40
    assert "content" not in first["version"] and "pages" not in first["version"]
    second = (await client.get(path + "?offset=20&limit=20&revision=1")).json()
    assert not ({x["id"] for x in first["chunks"]} & {x["id"] for x in second["chunks"]})
    assert (await client.get(path + "?limit=101")).status_code == 422
    source = (await client.get(path + "/source")).json()
    assert len(source["text"]) == 10000 and source["total"] > 10000
    policy = first["version"]["chunking"]
    body = {
        "revision": 1,
        "policy": policy,
        "offset": 20,
        "replace_count": 20,
        "chunks": [{k: x[k] for k in ("content", "source", "enabled")} for x in second["chunks"]],
    }
    body["chunks"][0]["content"] = "Edited second page"
    assert (await client.put(path + "/draft", json=body)).status_code == 200
    after = (await client.get(path)).json()
    assert after["total"] == first["total"]
    assert [x["content"] for x in after["chunks"]] == [x["content"] for x in first["chunks"]]
    assert (await client.get(path + "?offset=20")).json()["chunks"][0][
        "content"
    ] == "Edited second page"
    assert (await client.get(path + "?revision=1")).status_code == 409
    assert (await client.put(path + "/draft", json=body)).status_code == 409
    preview = (
        await client.post(path + "/preview?offset=20&limit=10&revision=2", json=policy)
    ).json()
    assert len(preview["chunks"]) == 10 and preview["total"] > 40
    assert (await client.get(path + "?offset=10000")).json()["chunks"] == []


def test_office_and_table_extraction():
    doc = Document()
    doc.add_heading("退款规则", level=1)
    doc.add_paragraph("收到商品七天内可以申请退款。")
    out = io.BytesIO()
    doc.save(out)
    pages = parse_file("policy.docx", out.getvalue())
    assert "退款规则" in pages[0]["text"]
    book = Workbook()
    sheet = book.active
    sheet.append(["问题", "答案"])
    sheet.append(["如何退款", "在订单页面申请。" * 60])
    out = io.BytesIO()
    book.save(out)
    pages = parse_file("faq.xlsx", out.getvalue())
    chunks = chunk_pages(pages, ChunkingPolicy(strategy="faq", size=200, overlap=20))
    assert len(chunks) > 1
    assert all(
        c["content"].startswith("问题：如何退款") and len(c["content"]) <= 200 for c in chunks
    )
    assert chunks[0]["source"]["sheet"] == "Sheet"
    assert chunk_pages(pages, ChunkingPolicy(strategy="table", sheet_name="Missing")) == []


def test_rank_fusion_does_not_compare_raw_scores():
    a = {"chunk_id": "a", "score": 0.01}
    b = {"chunk_id": "b", "score": 1000}
    hits = fuse([("keyword", [a, b]), ("vector", [a])])
    assert hits[0]["chunk_id"] == "a"
    assert hits[0]["channels"] == ["keyword", "vector"]


def test_pdf_pptx_csv_and_whole_document_limits():
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    pdf = PdfWriter()
    page = pdf.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 200 Td (Refund within seven days) Tj ET")
    page[NameObject("/Contents")] = stream
    output = io.BytesIO()
    pdf.write(output)
    parsed = parse_file("policy.pdf", output.getvalue())
    assert "Refund" in parsed[0]["text"] and parsed[0]["source"]["page_num"] == 1
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Support policy"
    output = io.BytesIO()
    slides.save(output)
    assert "Support policy" in parse_file("policy.pptx", output.getvalue())[0]["text"]
    csv = parse_file("faq.csv", "标题,\n问题,答案\n退款,七天内申请\n".encode())
    chunks = chunk_pages(csv, ChunkingPolicy(strategy="faq", header_row=2))
    assert "问题：退款" in chunks[0]["content"] and chunks[0]["source"]["row_start"] == 3
    with pytest.raises(ValueError, match="whole_document_exceeds"):
        chunk_pages(
            [{"page_num": 1, "text": "a" * 70}, {"page_num": 2, "text": "b" * 70}],
            ChunkingPolicy(strategy="whole", size=100, overlap=0),
        )


@pytest.mark.asyncio
@pytest.mark.postgres
async def test_workspace_publish_draft_isolation_and_tenant(platform):
    client, services, settings = platform
    w = services.platform.knowledge_workspace
    kb = (await client.post("/api/v1/knowledge-bases", json={"name": "Workspace"})).json()
    state = (await client.get(ROOT + f"/libraries/{kb['id']}")).json()
    assert state["vector_ready"] is False and state["published_document_count"] == 0
    unavailable = await client.post(
        ROOT + "/search", json={"query": "退款", "kb_ids": [kb["id"]], "mode": "hybrid"}
    )
    assert unavailable.status_code == 409
    assert "knowledge_index_not_ready" in unavailable.text
    upload = await client.post(
        ROOT + "/uploads?filename=policy.md", content="# 退款\n七天内申请退款。".encode()
    )
    assert upload.status_code == 202, upload.text
    did = upload.json()["id"]
    assert await w.run_next(settings.tenant_id)
    detail = (await client.get(ROOT + f"/documents/{did}")).json()
    assert detail["processing"] == "ready", detail
    assert detail["chunks"]
    listed = (await client.get(ROOT + "/documents")).json()["items"][0]
    assert listed["chunk_count"] == detail["total"]
    assert listed["published_library_count"] == listed["pending_library_count"] == 0
    assert (await client.get(ROOT + "/documents")).json()["total"] == 1
    assert (
        await client.post(ROOT + f"/libraries/{kb['id']}/documents", json={"document_id": did})
    ).status_code == 200
    assert (await client.post(ROOT + f"/libraries/{kb['id']}/publish")).status_code == 202
    assert await w.run_next(settings.tenant_id)
    result = await client.post(ROOT + "/search", json={"query": "退款", "kb_ids": [kb["id"]]})
    assert result.status_code == 200, result.text
    assert result.json()["items"][0]["version"] == 1
    edited = await client.put(
        ROOT + f"/documents/{did}/draft",
        json={"revision": 1, "chunks": [{"content": "一年保修。"}]},
    )
    assert edited.status_code == 200, edited.text
    listed = (await client.get(ROOT + "/documents")).json()["items"][0]
    assert (
        listed["library_count"]
        == listed["published_library_count"]
        == listed["pending_library_count"]
        == 1
    )
    result = await client.post(ROOT + "/search", json={"query": "退款", "kb_ids": [kb["id"]]})
    assert result.json()["items"][0]["version"] == 1
    stale = await client.put(ROOT + f"/documents/{did}/draft", json={"revision": 1})
    assert stale.status_code == 409
    with pytest.raises(Exception, match="not_found"):
        await w.detail("other-tenant", UUID(did))
    assert (
        await client.patch(ROOT + f"/documents/{did}", json={"enabled": False})
    ).status_code == 200
    result = await client.post(ROOT + "/search", json={"query": "退款", "kb_ids": [kb["id"]]})
    assert result.json()["items"] == []


@pytest.mark.asyncio
@pytest.mark.postgres
async def test_workspace_cancel_and_viewer(platform):
    client, services, settings = platform
    uploaded = (await client.post(ROOT + "/uploads?filename=a.txt", content=b"hello world")).json()
    assert (await client.post(ROOT + f"/tasks/{uploaded['task_id']}/cancel")).status_code == 200
    assert not await services.platform.knowledge_workspace.run_next(settings.tenant_id)
    cancelled = (await client.get(ROOT + f"/documents/{uploaded['id']}")).json()
    assert cancelled["processing"] == "cancelled"
    denied = await client.post(
        ROOT + "/folders",
        json={"name": "forbidden"},
        headers={"Authorization": "Bearer platform-viewer"},
    )
    assert denied.status_code == 403


@pytest.mark.asyncio
@pytest.mark.postgres
@pytest.mark.skipif(os.environ.get("KNOWLEDGE_LIVE") != "1", reason="Opt-in real Milvus test")
async def test_live_vector_publication_and_agent_retrieval(platform):
    from test_model_gateway import setup

    client, services, settings = platform
    _, _, embedding = await setup(client, "embed")
    _, _, rerank = await setup(client, "rerank")
    kid = (await client.post("/api/v1/knowledge-bases", json={"name": "Milvus live"})).json()["id"]
    settings_response = await client.put(
        ROOT + f"/libraries/{kid}/settings",
        json={
            "revision": 1,
            "settings": {
                "mode": "hybrid",
                "embedding": {"id": embedding["id"], "version": 1},
                "rerank": {"id": rerank["id"], "version": 1},
            },
        },
    )
    assert settings_response.status_code == 200, settings_response.text
    upload = await client.post(
        ROOT + f"/uploads?filename=live.md&mode=automatic&kb={kid}",
        content="# 退款政策\n签收七天内可申请退款。".encode(),
    )
    assert upload.status_code == 202, upload.text
    w = services.platform.knowledge_workspace
    await w.run_next(settings.tenant_id)
    await w.run_next(settings.tenant_id)
    tasks = (await client.get(ROOT + "/tasks")).json()
    assert all(t["status"] == "completed" for t in tasks), tasks
    response = await client.post(ROOT + "/search", json={"query": "退款", "kb_ids": [kid]})
    assert response.status_code == 200, response.text
    assert response.json()["items"][0]["channels"] == ["keyword", "vector"]
    assert "rerank_score" in response.json()["items"][0]
    assert await services.platform.retrieval.for_agent(settings.tenant_id, "退款", [UUID(kid)])
    did = upload.json()["id"]
    replacement = await client.post(
        ROOT + f"/documents/{did}/replace?filename=live.md",
        content="# 保修\n设备支持一年保修。".encode(),
    )
    assert replacement.status_code == 202, replacement.text
    await w.run_next(settings.tenant_id)
    await client.post(ROOT + f"/libraries/{kid}/documents", json={"document_id": did})
    await client.post(ROOT + f"/libraries/{kid}/publish")
    original_request = services.platform.vector_index.request
    from agent_platform.platform.persistence.store import DomainError

    async def fail_upsert(path, body):
        if path == "entities/upsert":
            raise DomainError("test_index_provider_failure", 502)
        return await original_request(path, body)

    services.platform.vector_index.request = fail_upsert
    try:
        await w.run_next(settings.tenant_id)
    finally:
        services.platform.vector_index.request = original_request
    library = (await client.get(ROOT + f"/libraries/{kid}")).json()
    assert library["documents"][0]["published_version"] == 1
    response = await client.post(ROOT + "/search", json={"query": "退款", "kb_ids": [kid]})
    assert response.status_code == 200 and response.json()["items"][0]["version"] == 1
    # Dispose only the collection created by this test.
    from agent_platform.platform.persistence.store import one

    async with services.platform.engine.connect() as c:
        index = await one(
            c, "SELECT collection_name FROM knowledge_vector_indexes WHERE kb_id=:k", k=UUID(kid)
        )
    await services.platform.vector_index.request(
        "collections/drop", {"collectionName": index["collection_name"]}
    )
