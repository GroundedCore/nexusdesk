"""Local file extraction. OCR and legacy conversion remain external adapters."""

import csv
import io
import zipfile
from pathlib import PurePath

from agent_platform.platform.persistence.store import DomainError

LOCAL_FORMATS = {".txt", ".md", ".pdf", ".docx", ".pptx", ".xlsx", ".csv"}


def parse_file(filename: str, content: bytes, max_chars: int = 2_000_000) -> list[dict]:
    extension = PurePath(filename).suffix.lower()
    if extension not in LOCAL_FORMATS:
        raise DomainError("external_parser_required", 422)
    source = io.BytesIO(content)
    if extension in {".docx", ".pptx", ".xlsx"}:
        with zipfile.ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100_000_000:
                raise DomainError("expanded_document_too_large", 413)
        source.seek(0)
    pages = []

    def page(text, elements=None):
        number = len(pages) + 1
        location = {"page_num": number} if extension in {".pdf", ".pptx"} else {"block_num": number}
        pages.append(
            {"page_num": number, "source": location, "text": text, "elements": elements or []}
        )

    def table(rows, sheet):
        matrix, total = [], 0
        for index, row in enumerate(rows):
            converted = [str(v) if v is not None else "" for v in row]
            total += sum(len(v) for v in converted)
            if total > max_chars or index >= 10000:
                raise DomainError("extracted_table_too_large", 413)
            matrix.append(converted)
        rows = matrix
        if rows and any(any(row) for row in rows):
            page(
                "\n".join(" | ".join(row) for row in rows),
                [{"type": "table", "sheet": sheet, "header": rows[0], "rows": rows[1:]}],
            )

    if extension in {".txt", ".md"}:
        page(content.decode("utf-8-sig"))
    elif extension == ".csv":
        table(csv.reader(io.StringIO(content.decode("utf-8-sig"))), filename)
    elif extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(source)
        if reader.is_encrypted:
            raise DomainError("encrypted_document", 422)
        for item in reader.pages:
            text = item.extract_text() or ""
            if not text.strip():
                raise DomainError("ocr_required", 422)
            page(text)
    elif extension == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(source)
        paragraphs = []
        for item in doc.iter_inner_content():
            if isinstance(item, Paragraph) and item.text.strip():
                style = item.style.name if item.style else ""
                prefix = ""
                if style.startswith("Heading ") and style[-1:].isdigit():
                    prefix = "#" * min(6, int(style[-1])) + " "
                if prefix and paragraphs:
                    page("\n\n".join(paragraphs), [{"type": "paragraph"}])
                    paragraphs = []
                paragraphs.append(prefix + item.text)
            elif isinstance(item, Table):
                if paragraphs:
                    page("\n\n".join(paragraphs), [{"type": "paragraph"}])
                    paragraphs = []
                table([[cell.text for cell in row.cells] for row in item.rows], "")
        if paragraphs:
            page("\n\n".join(paragraphs), [{"type": "paragraph"}])
    elif extension == ".pptx":
        from pptx import Presentation

        for slide in Presentation(source).slides:
            text = []
            elements = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text.append(shape.text)
                if shape.has_table:
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    text.extend(" | ".join(row) for row in rows)
                    elements.append({"type": "table", "header": rows[0], "rows": rows[1:]})
            page("\n\n".join(text), elements)
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(source, read_only=True, data_only=True)
        try:
            for sheet in book:
                table(sheet.iter_rows(values_only=True), sheet.title)
        finally:
            book.close()
    if not pages or not any(p["text"].strip() for p in pages):
        raise DomainError("document_has_no_text", 422)
    if sum(len(p["text"]) for p in pages) > max_chars:
        raise DomainError("extracted_text_too_large", 413)
    return pages
